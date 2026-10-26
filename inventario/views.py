from datetime import date
import openpyxl
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .forms import EnvioViveroForm, InventarioSobranteForm, VentaMenorForm
from .models import EnvioVivero, InventarioSobrante, Planta, VentaMenor


# --- VISTAS DE COMPUTADORA ---
def panel_computadora(request):
  ventas = VentaMenor.objects.all().order_by('-id')
  envios = EnvioVivero.objects.all().order_by('-id')

  form_venta = VentaMenorForm(prefix='venta')
  form_envio = EnvioViveroForm(prefix='envio')

  if request.method == 'POST':
    if 'guardar_venta' in request.POST:
      form_venta = VentaMenorForm(request.POST, prefix='venta')
      if form_venta.is_valid():
        form_venta.save()
        return redirect('panel_pc')
    elif 'guardar_envio' in request.POST:
      form_envio = EnvioViveroForm(request.POST, prefix='envio')
      if form_envio.is_valid():
        form_envio.save()
        return redirect('panel_pc')

  return render(
      request,
      'inventario/panel_pc.html',
      {
          'form_venta': form_venta,
          'form_envio': form_envio,
          'ventas': ventas,
          'envios': envios,
      },
  )


def editar_venta(request, pk):
  venta = get_object_or_404(VentaMenor, pk=pk)
  if request.method == 'POST':
    form = VentaMenorForm(request.POST, instance=venta)
    if form.is_valid():
      form.save()
      return redirect('panel_pc')
  else:
    form = VentaMenorForm(instance=venta)
  return render(
      request, 'inventario/editar_venta.html', {'form': form, 'venta': venta}
  )


def eliminar_venta(request, pk):
  venta = get_object_or_404(VentaMenor, pk=pk)
  venta.delete()
  return redirect('panel_pc')


def editar_envio(request, pk):
  envio = get_object_or_404(EnvioVivero, pk=pk)
  if request.method == 'POST':
    form = EnvioViveroForm(request.POST, instance=envio)
    if form.is_valid():
      form.save()
      return redirect('panel_pc')
  else:
    form = EnvioViveroForm(instance=envio)
  return render(
      request, 'inventario/editar_envio.html', {'form': form, 'envio': envio}
  )


def eliminar_envio(request, pk):
  envio = get_object_or_404(EnvioVivero, pk=pk)
  envio.delete()
  return redirect('panel_pc')


# --- VISTAS DE CELULAR ---
def panel_celular(request):
  sobrantes = InventarioSobrante.objects.all().order_by('-id')
  if request.method == 'POST':
    form = InventarioSobranteForm(request.POST)
    if form.is_valid():
      form.save()
      return redirect('panel_celular')
  else:
    form = InventarioSobranteForm()
  return render(
      request,
      'inventario/panel_celular.html',
      {'form': form, 'sobrantes': sobrantes},
  )


def editar_sobrante(request, pk):
  sobrante = get_object_or_404(InventarioSobrante, pk=pk)
  if request.method == 'POST':
    form = InventarioSobranteForm(request.POST, instance=sobrante)
    if form.is_valid():
      form.save()
      return redirect('panel_celular')
  else:
    form = InventarioSobranteForm(instance=sobrante)
  return render(
      request,
      'inventario/editar_sobrante.html',
      {'form': form, 'sobrante': sobrante},
  )  # <- Quitale la 's' aquí si tu archivo no la lleva


def eliminar_sobrante(request, pk):
  sobrante = get_object_or_404(InventarioSobrante, pk=pk)
  sobrante.delete()
  return redirect('panel_celular')


# --- EXPORTAR EXCEL CELULAR ---
def exportar_excel_celular(request):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Inventario Celular'
  ws.views.sheetView[0].showGridLines = True

  font_title = Font(name='Calibri', size=11, bold=True)
  font_header = Font(name='Calibri', size=10, bold=True, color='000000')

  fill_green_light = PatternFill(
      start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'
  )
  fill_yellow = PatternFill(
      start_color='FFFF99', end_color='FFFF99', fill_type='solid'
  )
  fill_header_gray = PatternFill(
      start_color='D9D9D9', end_color='D9D9D9', fill_type='solid'
  )

  thin_border = Border(
      left=Side(style='thin', color='000000'),
      right=Side(style='thin', color='000000'),
      top=Side(style='thin', color='000000'),
      bottom=Side(style='thin', color='000000'),
  )

  ws.merge_cells('A1:C1')
  ws['A1'] = 'FECHA'
  ws['A1'].font = font_title
  ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
  ws['A1'].fill = fill_header_gray
  for col in range(1, 4):
    ws.cell(row=1, column=col).border = thin_border

  ws.merge_cells('D1:G1')
  fecha_hoy = date.today().strftime('%A, %d de %B de %Y')
  ws['D1'] = fecha_hoy
  ws['D1'].font = font_title
  ws['D1'].alignment = Alignment(horizontal='center', vertical='center')
  ws['D1'].fill = PatternFill(
      start_color='FCE4D6', end_color='FCE4D6', fill_type='solid'
  )
  for col in range(4, 8):
    ws.cell(row=1, column=col).border = thin_border

  headers = [
      'PLANTA',
      'BANDEJA',
      'FILA',
      'BANDEJAS',
      'plantas x bandeja',
      'FILAS',
      'plantas en filas',
      'SUMA TOTAL',
  ]
  ws.append([])
  ws.append(headers)

  for col_num, header in enumerate(headers, start=1):
    cell = ws.cell(row=3, column=col_num)
    cell.font = font_header
    cell.alignment = Alignment(
        horizontal='center', vertical='center', wrap_text=True
    )
    cell.border = thin_border
    if col_num in [1, 2, 3]:
      cell.fill = fill_header_gray
    else:
      cell.fill = fill_green_light

  plantas = Planta.objects.all().order_by('nombre')

  for row_idx, planta in enumerate(plantas, start=4):
    sobrante = (
        InventarioSobrante.objects.filter(planta=planta).order_by('-id').first()
    )

    bandejas_completas = sobrante.bandejas_completas if sobrante else 0
    filas_sueltas = sobrante.filas_sueltas if sobrante else 0

    val_bandejas = bandejas_completas if bandejas_completas > 0 else ''
    plantas_x_bandeja = (
        bandejas_completas * planta.capacidad_bandeja
        if bandejas_completas > 0
        else 0
    )
    val_filas = filas_sueltas if filas_sueltas > 0 else ''
    plantas_en_filas = filas_sueltas * planta.plantas_por_fila
    suma_total = plantas_x_bandeja + plantas_en_filas

    row_data = [
        planta.nombre,
        planta.capacidad_bandeja,
        planta.plantas_por_fila,
        val_bandejas,
        plantas_x_bandeja,
        val_filas,
        plantas_en_filas,
        suma_total,
    ]

    ws.append(row_data)

    for col_num in range(1, 9):
      cell = ws.cell(row=row_idx, column=col_num)
      cell.border = thin_border
      cell.font = Font(name='Calibri', size=10)

      if col_num == 1:
        cell.alignment = Alignment(horizontal='left', vertical='center')
      else:
        cell.alignment = Alignment(horizontal='center', vertical='center')

      if col_num == 8:
        cell.fill = fill_yellow
        cell.font = Font(name='Calibri', size=10, bold=True)

  for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

  response = HttpResponse(
      content_type=(
          'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
      )
  )
  response['Content-Disposition'] = (
      'attachment; filename=reporte_celular_sobrantes.xlsx'
  )
  wb.save(response)
  return response


# --- EXPORTAR EXCEL PC ---
def exportar_excel_pc(request):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Control PC'
  ws.views.sheetView[0].showGridLines = True

  font_title = Font(name='Calibri', size=11, bold=True)
  font_header = Font(name='Calibri', size=10, bold=True, color='000000')

  fill_green_light = PatternFill(
      start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'
  )
  fill_pink = PatternFill(
      start_color='F8CBAD', end_color='F8CBAD', fill_type='solid'
  )
  fill_blue_light = PatternFill(
      start_color='DDEBF7', end_color='DDEBF7', fill_type='solid'
  )
  fill_header_gray = PatternFill(
      start_color='D9D9D9', end_color='D9D9D9', fill_type='solid'
  )

  thin_border = Border(
      left=Side(style='thin', color='000000'),
      right=Side(style='thin', color='000000'),
      top=Side(style='thin', color='000000'),
      bottom=Side(style='thin', color='000000'),
  )

  # Fila 1: Fecha combinada
  ws.merge_cells('A1:F1')
  ws['A1'] = 'FECHA'
  ws['A1'].font = font_title
  ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
  ws['A1'].fill = fill_header_gray

  ws.merge_cells('G1:J1')
  fecha_hoy = date.today().strftime('%A, %d de %B de %Y')
  ws['G1'] = fecha_hoy
  ws['G1'].font = font_title
  ws['G1'].alignment = Alignment(horizontal='center', vertical='center')
  ws['G1'].fill = PatternFill(
      start_color='FCE4D6', end_color='FCE4D6', fill_type='solid'
  )

  for col in range(1, 11):
    ws.cell(row=1, column=col).border = thin_border

  # Fila 3: Encabezados principales de bloques
  ws.merge_cells('A3:B3')
  ws['A3'] = 'TOTAL ENVIOS'
  ws['A3'].font = font_header
  ws['A3'].fill = fill_green_light
  ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
  ws['A3'].border = thin_border
  ws['B3'].border = thin_border

  ws.merge_cells('C3:E3')
  ws['C3'] = 'CUADRO DE ENVIOS'
  ws['C3'].font = font_header
  ws['C3'].fill = fill_green_light
  ws['C3'].alignment = Alignment(horizontal='center', vertical='center')
  for c in range(3, 6):
    ws.cell(row=3, column=c).border = thin_border

  ws.merge_cells('G3:H3')
  ws['G3'] = 'TOTAL DE VENTAS'
  ws['G3'].font = font_header
  ws['G3'].fill = fill_pink
  ws['G3'].alignment = Alignment(horizontal='center', vertical='center')
  for c in range(7, 9):
    ws.cell(row=3, column=c).border = thin_border

  ws.merge_cells('I3:J3')
  ws['I3'] = 'CUADRO DE VENTAS'
  ws['I3'].font = font_header
  ws['I3'].fill = fill_blue_light
  ws['I3'].alignment = Alignment(horizontal='center', vertical='center')
  for c in range(9, 11):
    ws.cell(row=3, column=c).border = thin_border

  # Fila 4: Sub-encabezados de columnas
  sub_headers_vals = [
      ('PLANTA', fill_green_light),
      ('TOTAL', fill_green_light),
      ('PLANTA', fill_green_light),
      ('BANDE', fill_green_light),
      ('TOTAL', fill_green_light),
      ('', None),
      ('PLANTA', fill_pink),
      ('CANTIDAD', fill_pink),
      ('PLANTA', fill_blue_light),
      ('CANTIDAD', fill_blue_light),
  ]

  for c_idx, (val, fill_color) in enumerate(sub_headers_vals, start=1):
    cell = ws.cell(row=4, column=c_idx, value=val)
    cell.font = Font(name='Calibri', size=9, bold=True)
    cell.border = thin_border
    if fill_color:
      cell.fill = fill_color
    if c_idx != 6:
      cell.alignment = Alignment(horizontal='center', vertical='center')

  plantas = Planta.objects.all().order_by('nombre')
  envios_list = list(EnvioVivero.objects.all().order_by('id'))
  ventas_list = list(VentaMenor.objects.all().order_by('id'))

  max_rows = max(len(plantas), len(envios_list), len(ventas_list))

  # Rellenar filas de manera explícita para evitar errores de celdas vacías
  for i in range(max_rows):
    row_idx = 5 + i
    p_name = plantas[i].nombre if i < len(plantas) else ''
    total_env_planta = (
        sum(
            e.total_plantas
            for e in EnvioVivero.objects.filter(planta=plantas[i])
        )
        if i < len(plantas)
        else ''
    )

    env_planta = envios_list[i].planta.nombre if i < len(envios_list) else ''
    env_bandeja = envios_list[i].bandejas if i < len(envios_list) else ''
    env_total = envios_list[i].total_plantas if i < len(envios_list) else ''

    total_vent_planta = (
        sum(
            v.cantidad_vendida
            for v in VentaMenor.objects.filter(planta=plantas[i])
        )
        if i < len(plantas)
        else ''
    )

    vent_planta = ventas_list[i].planta.nombre if i < len(ventas_list) else ''
    vent_cant = ventas_list[i].cantidad_vendida if i < len(ventas_list) else ''

    row_values = [
        p_name,
        total_env_planta,
        env_planta,
        env_bandeja,
        env_total,
        '',
        p_name,
        total_vent_planta,
        vent_planta,
        vent_cant,
    ]

    for c_idx, val in enumerate(row_values, start=1):
      cell = ws.cell(row=row_idx, column=c_idx, value=val)
      cell.border = thin_border
      cell.font = Font(name='Calibri', size=9)

      # Alinear texto a la izquierda en plantas, al centro en números
      if c_idx in [1, 3, 7, 9]:
        cell.alignment = Alignment(horizontal='left', vertical='center')
      else:
        cell.alignment = Alignment(horizontal='center', vertical='center')

      # Aplicar color de fondo por bloques de manera estricta
      if c_idx in [1, 2, 3, 4, 5]:
        cell.fill = fill_green_light
      elif c_idx in [7, 8]:
        cell.fill = fill_pink
      elif c_idx in [9, 10]:
        cell.fill = fill_blue_light

  for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 2, 12)

  response = HttpResponse(
      content_type=(
          'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
      )
  )
  response['Content-Disposition'] = (
      'attachment; filename=reporte_pc_vivero.xlsx'
  )
  wb.save(response)
  return response